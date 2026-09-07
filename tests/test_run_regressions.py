"""Regression coverage for the September 5 run. No network or credentials."""
from __future__ import annotations

import contextlib
import copy
import importlib.util
import io
import json
import os
import subprocess
import sys
import tempfile
import types
import unittest
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "claude/app/mp_v01/src"))
import app_paths
import excel_report as ex
import generate_picks as gp
from adapters import massive_options as mv
from gates.risk import RiskLimits
from backtest.costs import CostModel
from strategy.picks import ExitPolicy, freeze, generate_picks, select_contract, verify
from strategy.variants import BY_NAME

spec = importlib.util.spec_from_file_location("fetch_data", ROOT / "claude/app/mp_v01/fetch_data.py")
fetch = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fetch)

try:
    import openpyxl
except ImportError:
    openpyxl = None


def frozen_run(picks):
    return freeze("2026-09-04", picks, exit_policy=ExitPolicy(), universe=["SPY"],
                  generated_utc="2026-09-04T19:00:00+00:00", source_files={})


class RunRegressions(unittest.TestCase):
    def test_expiration_coverage_is_a_duration_not_six_entries(self):
        today = date(2026, 9, 5)
        expiries = [(today + timedelta(days=d)).isoformat() for d in (1, 2, 3, 4, 5, 6, 21, 35, 60, 61)]
        result = fetch.select_expiries(list(reversed(expiries)) + ["bad", None], today)
        self.assertEqual(result, expiries[:-1])
        self.assertEqual(sum(21 <= (date.fromisoformat(e) - today).days <= 60
                             for e in result), 3)

    def test_fetch_reaches_the_entry_band_and_survives_one_failed_expiry(self):
        today = datetime.now(fetch.US_EASTERN).date()
        expiries = [(today + timedelta(days=d)).isoformat() for d in (1, 2, 3, 4, 5, 6, 35)]
        calls = []
        quote = {"strike": 100, "bid": 4, "ask": 4.1, "volume": 1000,
                 "openInterest": 5000, "impliedVolatility": .3}
        frame = types.SimpleNamespace(iterrows=lambda: iter([(0, quote)]))

        class Ticker:
            options = expiries

            def option_chain(self, expiry):
                calls.append(expiry)
                if expiry == expiries[1]:
                    raise RuntimeError("one failed response")
                return types.SimpleNamespace(calls=frame, puts=frame)

        with tempfile.TemporaryDirectory() as tmp, patch.object(fetch, "DATA_DIR", tmp), \
                patch.dict(sys.modules, {"yfinance": types.SimpleNamespace(Ticker=lambda _: Ticker())}), \
                contextlib.redirect_stdout(io.StringIO()):
            before = datetime.now(timezone.utc)
            self.assertGreater(fetch.snapshot_chains(["SPY"]), 0)
            doc = json.loads(next((Path(tmp) / "chains").glob("*.json")).read_text())
            self.assertIn(expiries[-1], calls)
            self.assertEqual(doc["failed_expiries"], [expiries[1]])
            self.assertGreater(doc["contracts_in_entry_band"], 0)
            self.assertGreaterEqual(datetime.fromisoformat(doc["available_time"]), before)

    def test_empty_fetch_does_not_hide_the_previous_snapshot(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(fetch, "DATA_DIR", tmp), \
                patch.dict(sys.modules, {"yfinance": types.SimpleNamespace(
                    Ticker=lambda _: types.SimpleNamespace(options=[]))}), \
                contextlib.redirect_stdout(io.StringIO()):
            chains = Path(tmp) / "chains"
            chains.mkdir()
            old = chains / "SPY__v20260904T190000Z.json"
            old.write_text('{"old": true}')
            self.assertEqual(fetch.snapshot_chains(["SPY"]), 0)
            self.assertEqual(list(chains.iterdir()), [old])

    def test_contract_rejection_names_dte_failure_without_relaxing_limits(self):
        rows = ex.build_option_rows("SPY", {
            "snapshot_time_utc": "2026-09-04T19:00:00+00:00",
            "contracts": [{"type": "CALL", "expiration": "2026-09-11", "strike": 100,
                           "bid": 4, "ask": 4.1, "mid": 4.05, "status": "OK",
                           "open_interest": 5000, "volume": 1000}],
        }, [{"date": "2026-09-03", "close": 100}], risk_free_rate=.04,
           limits=RiskLimits(), costs=CostModel())
        selected, reason = select_contract(rows, kind="CALL", target_abs_delta=.5)
        self.assertIsNone(selected)
        self.assertIn("DTE outside entry band: 1/1", reason)

    def test_cutoff_selects_the_earlier_available_snapshot(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / "chains"
            folder.mkdir()
            for stamp, available in (("190000", "19:00:00"), ("200000", "20:00:00")):
                (folder / f"SPY__v20260904T{stamp}Z.json").write_text(json.dumps({
                    "snapshot_time_utc": f"2026-09-04T{available}+00:00",
                    "available_time": f"2026-09-04T{available}+00:00"}))
            result = ex.find_chain_files(Path(tmp), datetime(2026, 9, 4, 19, 45, tzinfo=timezone.utc))
            self.assertIn("190000", result["SPY"].name)

    def test_weekend_and_stale_chains_produce_explicit_abstentions(self):
        with tempfile.TemporaryDirectory() as tmp:
            chain = Path(tmp) / "chain.json"
            chain.write_text(json.dumps({"snapshot_time_utc": "2026-09-03T19:00:00+00:00"}))
            data = {"rows": {"SPY": []}, "options": {"SPY": []}, "chain_files": {"SPY": str(chain)}}
            for day, word in (("2026-09-05", "weekend"), ("2026-09-04", "stale")):
                now = datetime.fromisoformat(day + "T19:00:00+00:00")
                payload = gp.prepare_inputs(data, day, now)
                picks = generate_picks(day, payload, variants=[BY_NAME["momentum"]], exit_policy=ExitPolicy())
                self.assertEqual(picks[0]["action"], "ABSTAIN")
                self.assertIn(word, picks[0]["reason"])

    def test_forward_records_cannot_be_backdated(self):
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as exc:
            gp.main(["--decision-date", "2020-01-01"])
        self.assertEqual(exc.exception.code, 2)

    def test_metadata_changes_are_detected_and_legacy_records_still_verify(self):
        original = frozen_run([])
        for key, value in (("decision_date", "2000-01-01"), ("generated_utc", "future"),
                           ("universe", ["OTHER"]), ("n_picks", 9), ("exit_policy", {})):
            edited = copy.deepcopy(original)
            edited[key] = value
            self.assertFalse(verify(edited), key)
        self.assertTrue(verify(original))
        legacy = dict(original, contract_version="0.1.0")
        legacy.pop("record_sha256")
        self.assertTrue(verify(legacy))
        new = dict(original)
        new.pop("record_sha256")
        self.assertFalse(verify(new))

    def test_massive_rejects_the_2012_response_seen_in_the_run(self):
        bad = {"status": "OK", "results": [{"ticker": "O:CYU121222C00060000",
               "underlying_ticker": "CYU", "expiration_date": "2012-12-22",
               "strike_price": 60, "contract_type": "call"}]}
        with patch.object(mv, "_get", return_value=bad):
            result = mv.probe("SPY", "2025-08-01")
            self.assertFalse(result["ok"])
            self.assertEqual(result["reason"], "INVALID_CONTRACT_RESPONSE")
            with self.assertRaises(mv.MassiveError):
                mv.list_contracts_as_of("SPY", "2025-08-01")

    def test_reference_probe_does_not_claim_price_entitlement(self):
        row = {"ticker": "O:SPY250919C00600000", "underlying_ticker": "SPY",
               "expiration_date": "2025-09-19", "strike_price": 600, "contract_type": "call"}
        with patch.object(mv, "_get", return_value={"status": "OK", "results": [row]}) as get:
            result = mv.probe("SPY", "2025-08-01")
            self.assertTrue(result["ok"])
            self.assertFalse(result["price_history_verified"])
            self.assertEqual(get.call_args.args[1]["expiration_date.gte"], "2025-08-01")
        with patch.object(mv, "_get", return_value={"status": "OK", "results": []}):
            self.assertFalse(mv.probe("SPY", "2025-08-01")["ok"])

    def test_massive_page_limit_does_not_silently_claim_complete_history(self):
        response = {"status": "OK", "results": [], "next_url": mv.BASE_URL + "/next"}
        with patch.object(mv, "_get", return_value=response) as get:
            with self.assertRaises(mv.MassiveError) as exc:
                list(mv._paginate(mv.CONTRACTS_PATH, {}, max_pages=1))
            self.assertEqual(exc.exception.kind, "INCOMPLETE_RESPONSE")
            self.assertEqual(get.call_count, 1)

    def test_output_root_and_migration_preserve_records_and_conflicts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            repo, runtime = root / "code", root / "outputs"
            (repo / "picks").mkdir(parents=True)
            (runtime / "picks").mkdir(parents=True)
            name = "picks_2026-09-04_1.json"
            original = json.dumps(frozen_run([])).encode()
            (repo / "picks" / name).write_bytes(original)
            (runtime / "picks" / name).write_text("different record")
            data = repo / "claude/app/mp_v01/data_store/bars"
            data.mkdir(parents=True)
            (data / "SPY__v1.json").write_text("{}")
            paths = app_paths.get_paths(runtime)
            first = app_paths.migrate_legacy(paths, repo=repo, settings={})
            self.assertEqual(first["copied"], 2)
            self.assertEqual(first["errors"], [])
            self.assertEqual((repo / "picks" / name).read_bytes(), original)
            self.assertEqual((runtime / "picks" / name).read_text(), "different record")
            copied = next((runtime / "picks").glob("*_legacy_*.json"))
            self.assertEqual(copied.read_bytes(), original)
            self.assertTrue(verify(json.loads(copied.read_text())))
            again = app_paths.migrate_legacy(paths, repo=repo, settings={})
            self.assertEqual(again["copied"], 0)
            self.assertTrue((paths.data / "bars/SPY__v1.json").is_file())
        with self.assertRaises(ValueError):
            app_paths.get_paths(ROOT / "outputs")

    def test_default_dashboard_and_nested_fetch_cli_work_from_another_directory(self):
        with tempfile.TemporaryDirectory() as tmp:
            runtime = Path(tmp) / "new" / "outputs"
            env = dict(os.environ, MONEYPRINTER_HOME=str(runtime))
            result = subprocess.run([sys.executable, str(ROOT / "dashboard.py")], cwd=tmp,
                                    env=env, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue((runtime / "dashboard.html").is_file())
            result = subprocess.run([sys.executable, str(ROOT / "claude/app/mp_v01/fetch_data.py"),
                                     "--help"], cwd=tmp, env=env, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("--data-dir", result.stdout)

    @unittest.skipIf(openpyxl is None, "openpyxl not installed")
    def test_valid_weekday_contract_flows_into_frozen_record_and_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data = root / "data"
            (data / "bars").mkdir(parents=True)
            (data / "chains").mkdir()
            (data / "bars/SPY__v20260904T190000Z.json").write_text(json.dumps({
                "rows": [{"date": "2026-09-03", "close": 100, "daily_total_return": .01}]}))
            (data / "chains/SPY__v20260904T190000Z.json").write_text(json.dumps({
                "snapshot_time_utc": "2026-09-04T19:00:00+00:00",
                "available_time": "2026-09-04T19:00:00+00:00",
                "contracts": [{"type": "CALL", "expiration": "2026-10-09", "strike": 100,
                               "bid": 4, "ask": 4.1, "mid": 4.05, "status": "OK",
                               "open_interest": 5000, "volume": 1000}]}))
            components = {"scaled": {k: .8 for k in BY_NAME["momentum"].weights},
                          "raw": {}, "last_available_date": "2026-09-03"}
            with patch.object(gp, "datetime", wraps=datetime) as clock, \
                    patch.object(gp.components, "compute", return_value=components), \
                    contextlib.redirect_stdout(io.StringIO()):
                clock.now.return_value = datetime(2026, 9, 4, 19, 30, tzinfo=timezone.utc)
                self.assertEqual(gp.main(["--data-dir", str(data), "--out-dir", str(root / "picks"),
                                          "--excel", str(root / "picks.xlsx"),
                                          "--variants", "momentum"]), 0)
            frozen = json.loads(next((root / "picks").glob("picks_*.json")).read_text())
            self.assertTrue(verify(frozen))
            self.assertEqual(frozen["n_picks"], 1)
            self.assertEqual(frozen["picks"][0]["gate_decision"], "PASS")
            wb = openpyxl.load_workbook(root / "picks.xlsx")
            try:
                self.assertEqual(wb["Pick_Runs"]["C2"].value, 1)
                self.assertEqual(wb["Pick_History"].max_row, 2)
                self.assertEqual(wb["Pick_Justifications"].max_row, 2)
            finally:
                wb.close()

    @unittest.skipIf(openpyxl is None, "openpyxl not installed")
    def test_zero_pick_workbook_opens_to_run_counts_and_keeps_all_abstentions(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            picks = root / "picks"
            picks.mkdir()
            decisions = [{"decision_date": "2026-09-04", "ticker": "SPY", "variant": f"v{i}",
                          "action": "ABSTAIN", "reason": "DTE outside entry band"} for i in range(20)]
            (picks / "picks_2026-09-04_1.json").write_text(json.dumps(frozen_run(decisions)))
            data = ex.collect(root / "store", picks_dir=picks)
            out = ex.write_workbook(data, root / "picks.xlsx", sections=ex.PICK_SECTIONS)
            wb = openpyxl.load_workbook(out)
            try:
                self.assertEqual(wb.active.title, "Pick_Runs")
                self.assertEqual(wb["Pick_Runs"]["C2"].value, 0)
                self.assertEqual(wb["Pick_Runs"]["D2"].value, 20)
                self.assertIn("NO PROPOSALS", wb["Pick_Runs"]["E2"].value)
                self.assertEqual(wb["Pick_Abstentions"].max_row, 21)
                self.assertEqual(wb["Pick_Abstentions"]["F2"].value, "OK")
            finally:
                wb.close()


class HistoricalScopeRegressions(unittest.TestCase):
    """PR19 changed the historical fetch and the diagnostic banner, and the whole
    suite still passed. These lock down the two contracts it broke silently."""

    def _params_for(self, **kwargs):
        sent = {}

        def capture(path, params, max_pages=20):
            sent.update(params)
            return []

        with patch.object(mv, "_paginate", capture):
            mv.list_contracts_as_of("SPY", "2024-01-02", **kwargs)
        return sent

    def test_point_in_time_fetch_asks_for_contracts_that_have_since_expired(self):
        """`expired` is evaluated against today, not as_of: the ladder's own OK
        condition is "includes at least one contract already expired today". So
        expired=false keeps only contracts still alive now, which is exactly the
        survivorship-biased sample. The default must stay expired=true."""
        self.assertEqual(self._params_for()["expired"], "true")
        self.assertEqual(self._params_for()["expiration_date.gte"], "2024-01-02")

    def test_fetch_cli_does_not_narrow_history_to_contracts_still_alive_today(self):
        """The regression PR19 shipped lived at the call site, not in the adapter
        default, so asserting the default alone does not cover it. Drive the real
        CLI path and check the scope it actually requests."""
        import fetch_massive as cli
        with tempfile.TemporaryDirectory() as tmp:
            args = types.SimpleNamespace(
                data_dir=tmp, tickers="SPY", as_of="2024-01-02", max_pages=2,
                bars=False, bars_from=None, max_contracts=1)
            with patch.object(mv, "list_contracts_as_of",
                              return_value=[]) as listed, \
                    contextlib.redirect_stdout(io.StringIO()):
                cli.do_fetch(args)
        self.assertEqual(listed.call_count, 1)
        self.assertNotIn(False, listed.call_args.kwargs.values())
        self.assertIs(listed.call_args.kwargs.get("expired", True), True)

    def test_expiry_floor_alone_does_not_make_the_sample_point_in_time(self):
        """The floor is satisfied by survivors too, so validation cannot catch an
        expired=false fetch. Demonstrate the loss the validator lets through."""
        today = date.today()
        chain = ["2024-01-19", "2024-06-21", "2025-01-17",
                 (today + timedelta(days=400)).isoformat()]

        def server(path, params, max_pages=20):
            want = params["expired"] == "true"
            floor = params["expiration_date.gte"]
            return [{"ticker": "O:SPY%sC00470000" % e.replace("-", ""),
                     "underlying_ticker": "SPY", "expiration_date": e,
                     "strike_price": 470.0, "contract_type": "call"}
                    for e in chain
                    if e >= floor and (date.fromisoformat(e) < today) == want]

        with patch.object(mv, "_paginate", server):
            historical = mv.list_contracts_as_of("SPY", "2024-01-02")
            survivors = mv.list_contracts_as_of("SPY", "2024-01-02", expired=False)

        # Both pass _validate_contracts, so nothing raises on the biased sample.
        self.assertEqual(len(historical), 3)
        self.assertEqual(len(survivors), 1)
        self.assertTrue(all(c["expiration"] >= "2024-01-02" for c in survivors))

    def test_production_fetch_scope_matches_the_canary_probe_scope(self):
        """The canary and the ladder both assert history with expired=true. If the
        production call site drifts from them, the probe stops describing the
        fetch it is meant to authorise."""
        row = {"ticker": "O:SPY250919C00600000", "underlying_ticker": "SPY",
               "expiration_date": "2025-09-19", "strike_price": 600,
               "contract_type": "call"}
        with patch.object(mv, "_get",
                          return_value={"status": "OK", "results": [row]}) as get:
            mv.probe("SPY", "2025-08-01")
            probe_params = get.call_args.args[1]
        self.assertEqual(probe_params["expired"], "true")
        self.assertEqual(self._params_for()["expired"], probe_params["expired"])

    def test_diagnostic_banner_and_verdict_column_stay_machine_readable(self):
        """The CLI summary line is a documented sentinel (CODE_REVIEW_PR11.md) and
        the verdict column is bracketed. PR19 shipped `DIAGNOSTI:` and a dangling
        `]` through a fully green matrix because nothing asserted the text."""
        import fetch_massive as cli
        steps = [{"step": "bare", "verdict": "OK", "asks": "answers?"},
                 {"step": "point in time", "verdict": "CONSISTENT",
                  "asks": "as_of effect?"}]
        args = types.SimpleNamespace(tickers="SPY", as_of="2024-01-02")
        out = io.StringIO()
        with patch.object(mv, "diagnose_access", return_value=steps), \
                contextlib.redirect_stdout(out):
            self.assertEqual(cli.do_diagnose(args), 0)
        text = out.getvalue()

        self.assertIn("DIAGNOSTIC: POINT_IN_TIME_CONSISTENT", text)
        self.assertNotIn("DIAGNOSTI:", text.replace("DIAGNOSTIC:", ""))
        for verdict in ("OK", "CONSISTENT"):
            self.assertIn("[%s" % verdict, text)
        for line in text.splitlines():
            self.assertEqual(line.count("["), line.count("]"), line)


if __name__ == "__main__":
    unittest.main(verbosity=2)
